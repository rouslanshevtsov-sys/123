#!/usr/bin/env python3
"""Итоговые артефакты: competitor_set.json + Excel (5 листов).

Использование:
    python build_outputs.py --context business_context.json --qualified qualified.json \
        --posts posts_er.json --excluded stage_filter.json --out-dir <dir> \
        [--manifest run_manifest.json]

Классификация типов конкурентов — по правилам references/classification.md
(поле "competitor_type" у квалифицированного кандидата; если пусто — indirect).
Формат Excel и колонки листа «Конкуренты» — references/excel-structure.md.
"""
import argparse
import json
import os
import sys
from datetime import datetime, timezone

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from context import get_criteria, load_context, sha256_file, make_manifest, save_json  # noqa: E402

COMPETITOR_COLUMNS = [
    "Название", "VK", "Сайт", "Тип конкурента", "Подписчики",
    "Последняя публикация", "Средний ER за месяц", "Ассортимент", "Аудитория",
    "География", "Ценовой сегмент", "Модель продаж", "Производство (если нужно)",
    "Онлайн/офлайн", "Почему конкурент", "Проверенные источники",
]


def vk_url(c):
    sn = c.get("screen_name")
    return f"https://vk.com/{sn}" if sn else f"https://vk.com/public{c.get('gid') or c.get('id')}"


def fmt_date(ts):
    if not ts:
        return ""
    return datetime.fromtimestamp(int(ts), tz=timezone.utc).strftime("%Y-%m-%d")


def build_competitor_row(c, posts):
    gid = str(c.get("gid") or c.get("id"))
    pr = posts.get(gid, {})
    return {
        "Название": c.get("name", ""),
        "VK": vk_url(c),
        "Сайт": c.get("site") or "",
        "Тип конкурента": c.get("competitor_type") or "indirect",
        "Подписчики": c.get("members_count"),
        "Последняя публикация": fmt_date(pr.get("last_publication_date")),
        "Средний ER за месяц": round(pr["avg_er"] * 100, 2) if pr.get("avg_er") is not None else "",
        "Ассортимент": c.get("assortment", ""),
        "Аудитория": c.get("audience", ""),
        "География": c.get("geography", ""),
        "Ценовой сегмент": c.get("price_segment", ""),
        "Модель продаж": c.get("sales_model", ""),
        "Производство (если нужно)": c.get("production", ""),
        "Онлайн/офлайн": c.get("format", ""),
        "Почему конкурент": c.get("why_competitor", ""),
        "Проверенные источники": "; ".join(c.get("sources_checked") or []),
    }


def write_excel(path, summary_rows, comp_rows, excluded_rows, post_rows, method_rows):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    wb = Workbook()

    def sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for r in rows:
            ws.append([r.get(h, "") for h in headers])
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(14, min(40, len(str(h)) + 6)))
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Сводка"
    sheet(ws, ["Параметр", "Значение"], summary_rows, {"Параметр": 40, "Значение": 70})
    ws2 = wb.create_sheet("Конкуренты")
    sheet(ws2, COMPETITOR_COLUMNS, comp_rows, {"Название": 36, "Почему конкурент": 50,
                                               "Проверенные источники": 45, "Ассортимент": 35})
    ws3 = wb.create_sheet("Проверены, но исключены")
    sheet(ws3, ["group_id", "Название", "Причина исключения"], excluded_rows,
          {"Название": 40, "Причина исключения": 60})
    ws4 = wb.create_sheet("Посты за месяц")
    sheet(ws4, ["group_id", "post_id", "owner_id", "Дата", "Просмотры", "Лайки",
                "Комментарии", "Репосты", "ER %", "Закреп", "Текст (превью)"], post_rows)
    ws5 = wb.create_sheet("Методика")
    sheet(ws5, ["Шаг", "Описание"], method_rows, {"Шаг": 30, "Описание": 100})
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--context", required=True)
    ap.add_argument("--qualified", required=True, help="квалифицированные конкуренты (JSON list)")
    ap.add_argument("--posts", required=True, help="posts_er.json")
    ap.add_argument("--excluded", required=True, help="stage_filter.json")
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--manifest", default=None)
    args = ap.parse_args()

    ctx, err = load_context(args.context)
    if err:
        print("STOP:", err)
        sys.exit(2)
    crit = get_criteria(ctx)
    with open(args.qualified, encoding="utf-8") as f:
        qualified = json.load(f)
    with open(args.posts, encoding="utf-8") as f:
        posts = json.load(f)
    with open(args.excluded, encoding="utf-8") as f:
        excluded = json.load(f)

    target = crit.get("count") or len(qualified)
    chosen = qualified[:target]

    comp_rows = [build_competitor_row(c, posts) for c in chosen]
    excluded_rows = [{"group_id": e.get("gid"), "Название": e.get("name"),
                      "Причина исключения": e.get("reason")} for e in excluded]
    post_rows = []
    for gid, pr in posts.items():
        for p in pr.get("scored_posts", []) + pr.get("excluded_posts", []):
            post_rows.append({
                "group_id": gid, "post_id": p.get("post_id"), "owner_id": p.get("owner_id"),
                "Дата": fmt_date(p.get("date")), "Просмотры": p.get("views"),
                "Лайки": p.get("likes"), "Комментарии": p.get("comments"),
                "Репосты": p.get("reposts"),
                "ER %": round(p["er"] * 100, 2) if p.get("er") is not None else "",
                "Закреп": "да" if p.get("is_pinned") else "",
                "Текст (превью)": p.get("text_preview") or p.get("excluded_reason", "")})

    ctx_hash = sha256_file(args.context)
    manifest_path = args.manifest or os.path.join(args.out_dir, "run_manifest.json")
    manifest = make_manifest(
        queries=crit.get("queries"),
        counters={"pool_size": crit.get("count") and None or None,
                  "qualified": len(qualified), "selected": len(chosen),
                  "excluded": len(excluded)},
        errors=[]) if not os.path.isfile(manifest_path) else json.load(open(manifest_path, encoding="utf-8"))
    manifest["business_context_sha256"] = ctx_hash
    manifest["counters"].update({"qualified": len(qualified), "selected": len(chosen),
                                 "excluded": len(excluded)})
    manifest["finished_at"] = datetime.now(timezone.utc).isoformat()
    manifest["shortfall"] = (None if len(chosen) >= (target or len(chosen)) else
                             {"target": target, "found": len(chosen),
                              "note": "criteria were NOT relaxed; see excluded reasons"})

    set_path = os.path.join(args.out_dir, "competitor_set.json")
    save_json({
        "status": "completed",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "business_context": {"path": os.path.basename(args.context),
                             "sha256": ctx_hash, "status": ctx.get("status")},
        "time_window_hours": 31 * 24,
        "target_count": target,
        "competitors": [{**{k: v for k, v in c.items()},
                         "vk_url": vk_url(c),
                         "avg_er": (posts.get(str(c.get('gid') or c.get('id')), {}) or {}).get("avg_er"),
                         "last_publication": (posts.get(str(c.get('gid') or c.get('id')), {}) or {}).get("last_publication_date")}
                        for c in chosen],
        "excluded_audit": excluded,
        "provenance": {"run_manifest": os.path.basename(manifest_path),
                       "inputs": [os.path.basename(args.qualified), os.path.basename(args.posts),
                                  os.path.basename(args.excluded)]},
    }, set_path)
    save_json(manifest, manifest_path)

    summary_rows = [
        {"Параметр": "Бизнес-контекст", "Значение": f"{os.path.basename(args.context)} (confirmed)"},
        {"Параметр": "Хэш контекста SHA-256", "Значение": ctx_hash},
        {"Параметр": "Целевое число конкурентов", "Значение": target},
        {"Параметр": "Фактически отобрано", "Значение": len(chosen)},
        {"Параметр": "Исключено кандидатов", "Значение": len(excluded)},
        {"Параметр": "Временное окно", "Значение": "31×24 ч"},
    ]
    method_rows = [
        {"Шаг": "Поиск", "Описание": "groups.search count=1000 по всем запросам, дедупликация по group_id"},
        {"Шаг": "Детализация", "Описание": "groups.getById пакетами"},
        {"Шаг": "Фильтрация", "Описание": "тип/доступность/подписчики/география/офтоп/включение-исключение карточки"},
        {"Шаг": "Публикации", "Описание": "wall.get count=100, offset-пагинация, окно 31×24 ч, закрепы отдельно"},
        {"Шаг": "ER", "Описание": "ER=(лайки+комменты+репосты)/просмотры, только views>0; нулевые — отдельно с причиной"},
        {"Шаг": "Классификация", "Описание": "direct/indirect/attention по правилам references/classification.md"},
        {"Шаг": "Проверка", "Описание": "сайты и первичные источники для каждого конкурента"},
    ]
    xlsx_path = os.path.join(args.out_dir, "competitors.xlsx")
    write_excel(xlsx_path, summary_rows, comp_rows, excluded_rows, post_rows, method_rows)
    print(f"[outputs] {set_path}")
    print(f"[outputs] {xlsx_path}")
    print(f"[outputs] {manifest_path}")
    print(f"[outputs] selected={len(chosen)} of target={target}, excluded={len(excluded)}")


if __name__ == "__main__":
    main()
