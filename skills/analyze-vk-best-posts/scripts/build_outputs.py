#!/usr/bin/env python3
"""Этап 4: итоговые файлы — XLSX (5 листов), vk_posts_dataset.json,
comments_replies.json (уже собран на этапе 1) и обезличенная медиасводка.

Требования к Excel (references/excel-structure.md):
листы Сводка / Лучшие посты / Клипы / Проверенные посты / Недоступно;
сортировка лучших постов: direct → indirect → attention, внутри по ER ↓;
кликабельные ссылки, ER в формате 0.00%, закреплённые заголовки, фильтры,
перенос текста, встроенные изображения/коллажи.

Использование: python build_outputs.py [--workdir DIR]
"""
import argparse
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (BEST_POSTS_COLUMNS, CLIPS_COLUMNS, EXCEL_SHEETS, MSK,  # noqa: E402
                    TYPE_ORDER, TYPE_RU, UNAVAILABLE_COLUMNS, VERIFIED_POSTS_COLUMNS,
                    load_json, path, save_json)

from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XImage  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HFONT = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="4472C4")


def style_sheet(ws, widths):
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for c in ws[1]:
        c.font = HFONT
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def link_cell(ws, r, col, url):
    c = ws.cell(row=r, column=col)
    if url:
        c.hyperlink = url
        c.font = Font(color="0563C1", underline="single")


def bucket(n):
    if n is None:
        return "н/д"
    for lim in (100_000, 200_000, 500_000, 1_000_000):
        if n < lim:
            return f"<{lim:,}"
    return ">=1,000,000"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workdir", default=None)
    args = ap.parse_args()
    if args.workdir:
        os.environ["VK_WORK_DIR"] = args.workdir

    meta = load_json("run_meta.json")
    posts = load_json("posts_analyzed.json")
    comments = load_json("comments_replies.json", default=[])
    groups = load_json("groups.json")
    unavailable = load_json("unavailable.json", default=[])
    clips = load_json("clips_processing.json", default=[])
    benchmarks = load_json("benchmarks.json", default={})

    best_posts = [p for p in posts if p["is_best"]]
    clip_by_post = {c["post_id"]: c for c in clips}

    # ---------------- vk_posts_dataset.json ----------------
    dataset = {
        "schema": "vk_posts_dataset/v1",
        "run": meta,
        "timezone": "Europe/Moscow",
        "window_hours": meta.get("window_hours", 168),
        "secrets_note": "Токены и секреты в файл не включаются",
        "communities": [
            {"name": g.get("resolved_name") or g["name"], "type": g["type"],
             "url": g.get("url"), "vk_id": g.get("vk_id"),
             "members_count": g.get("members_count")} for g in groups],
        "benchmarks": benchmarks,
        "posts": [],
        "clips_processing": clips,
        "unavailable": unavailable,
    }
    for p in posts:
        q = {k: v for k, v in p.items() if k not in ("videos", "clips", "clip_needs_video_get")}
        q["image_urls"] = [u for u, _ in p["images"] if u]
        cp = clip_by_post.get(p["post_id"])
        q["transcript"] = cp["transcript"] if cp else None
        q["transcript_status"] = (cp["status"] if cp else
                                  "not_applicable_no_confirmed_clip")
        dataset["posts"].append(q)
    save_json("vk_posts_dataset.json", dataset)

    # ---------------- обезличенная медиасводка ----------------
    anon = {"generated_at_msk": datetime.now(MSK).isoformat(),
            "window": {"start": meta["window_start_msk"], "end": meta["window_end_msk"],
                       "hours": meta.get("window_hours")},
            "communities": [], "per_community": []}
    for i, g in enumerate(groups):
        label = f"Конкурент {chr(65 + i) if i < 26 else i + 1}"
        anon["communities"].append({"label": label, "type": g["type"],
                                    "subscribers_band": bucket(g.get("members_count"))})
        gp = [p for p in posts if g.get("vk_id") is not None and p["owner_id"] == g["vk_id"]]
        gb = [p for p in gp if p["is_best"]]
        types = {}
        for p in gp:
            types[p["main_type"]] = types.get(p["main_type"], 0) + 1
        ers = sorted(x["er"] for x in gb if x["er"] is not None)
        anon["per_community"].append({
            "label": label, "posts_in_window": len(gp),
            "benchmark_er": benchmarks.get(str(g.get("vk_id"))),
            "best_posts": len(gb),
            "median_er_best": ers[len(ers) // 2] if ers else None,
            "post_types": types,
            "polls": sum(1 for p in gp if p["has_poll"]),
            "giveaways": sum(1 for p in gp if p["is_giveaway"]),
            "confirmed_clips": sum(1 for p in gp if p["main_type"] == "Клип"),
        })
    save_json("media_summary_anonymized.json", anon, indent=1)

    # ---------------- XLSX ----------------
    wb = Workbook()
    mem = {str(g["vk_id"]): g.get("members_count") for g in groups if g.get("vk_id")}
    snm = {str(g["vk_id"]): (g.get("screen_name") or "") for g in groups if g.get("vk_id")}

    ws = wb.active
    ws.title = EXCEL_SHEETS[0]
    rows = [["Параметр", "Значение"],
            ["Момент начала работы (Europe/Moscow)", meta["run_started_at_msk"]],
            ["Временное окно", f"{meta['window_start_msk']} — {meta['window_end_msk']} "
                               f"({meta.get('window_hours',168)} ч)"],
            ["Конкурентов в исходном JSON", len(groups)],
            ["Сообществ с данными", len({p['owner_id'] for p in posts})],
            ["Недоступно объектов", len(unavailable)],
            ["Всего проверенных публикаций", len(posts)],
            ["Без просмотров (ER не рассчитан)", sum(1 for p in posts if p["er"] is None)],
            ["Лучших постов (ER >= бенчмарка)", len(best_posts)],
            ["Подтверждённых VK Клипов", len(clips)],
            ["Строк комментариев/ответов", len(comments)],
            ["Типы постов", ", ".join(
                f"{t}: {sum(1 for p in posts if p['main_type'] == t)}"
                for t in ("Клип", "Видео", "Картинка", "Карусель", "Текстовый пост"))],
            ["Опросы (вложение poll)", sum(1 for p in posts if p["has_poll"])],
            ["Розыгрыши подтверждённые", sum(1 for p in posts if p["is_giveaway"])]]
    for r in rows:
        ws.append(r)
    style_sheet(ws, [45, 60])

    # Лучшие посты
    ws = wb.create_sheet(EXCEL_SHEETS[1])
    ws.append(BEST_POSTS_COLUMNS)
    bp_sorted = sorted(best_posts,
                       key=lambda p: (TYPE_ORDER.get(p["group_type"], 9),
                                      -(p["er"] if p["er"] is not None else -1)))
    for p in bp_sorted:
        cp = clip_by_post.get(p["post_id"])
        transcript = (cp["transcript"] if cp and cp["transcript"] else
                      ("—" if not cp else cp["status"]))
        ws.append([p["group_name"], TYPE_RU.get(p["group_type"], p["group_type"]),
                   p["url"], (p["text"] or "")[:3000], p["main_type"],
                   "Да" if p["has_poll"] else "", "Да" if p["is_giveaway"] else "",
                   p["likes"], p["comments"], p["views_for_er"], p["er"],
                   p["benchmark"], transcript, "", p["post_id"]])
        r = ws.max_row
        link_cell(ws, r, 3, p["url"])
        ws.cell(row=r, column=11).number_format = "0.00%"
        ws.cell(row=r, column=12).number_format = "0.00%"
        for col in (2, 4, 13, 14):
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        cf = p.get("collage_file")
        if cf and os.path.exists(cf):
            try:
                im = XImage(cf)
                scale = min(1.0, 300 / max(im.width, 1), 240 / max(im.height, 1))
                im.width = int(im.width * scale)
                im.height = int(im.height * scale)
                ws.add_image(im, f"N{r}")
                ws.row_dimensions[r].height = max(90, im.height * 0.75)
            except Exception as e:
                print(f"WARN xlsx image: {e}")
    style_sheet(ws, [26, 20, 30, 60, 12, 7, 9, 8, 11, 14, 10, 10, 16, 38, 14])

    # Клипы
    ws = wb.create_sheet(EXCEL_SHEETS[2])
    ws.append(CLIPS_COLUMNS)
    if not clips:
        ws.append(["—", "—", "—", "—", "—", "—",
                   "Подтверждённых VK Клипов за окно не найдено", "—"])
    for c in clips:
        ws.append([c["group_name"], c["post_url"], c["outer_type"], c["inner_type"],
                   c["clip_views"], c["file"], c["status"], c["transcript"]])
        r = ws.max_row
        link_cell(ws, r, 2, c["post_url"])
    style_sheet(ws, [26, 30, 12, 20, 14, 30, 30, 40])

    # Проверенные посты
    ws = wb.create_sheet(EXCEL_SHEETS[3])
    ws.append(VERIFIED_POSTS_COLUMNS)
    for p in sorted(posts, key=lambda x: (x["group_name"],
                                          -(x["er"] if x["er"] is not None else -1))):
        ws.append([p["group_name"], snm.get(str(p["owner_id"]), ""),
                   mem.get(str(p["owner_id"])), p["date_msk"], p["url"],
                   (p["text"] or "")[:3000], p["main_type"],
                   "Да" if p["has_poll"] else "", "Да" if p["is_giveaway"] else "",
                   p["likes"], p["comments"], p["reposts"], p["views_for_er"],
                   p["er"], p["benchmark"], p["result"], p["data_status"]])
        r = ws.max_row
        link_cell(ws, r, 5, p["url"])
        ws.cell(row=r, column=14).number_format = "0.00%"
        ws.cell(row=r, column=15).number_format = "0.00%"
        ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="top")
    style_sheet(ws, [26, 18, 11, 16, 30, 60, 12, 7, 9, 8, 11, 8, 14, 10, 10, 12, 25])

    # Недоступно
    ws = wb.create_sheet(EXCEL_SHEETS[4])
    ws.append(UNAVAILABLE_COLUMNS)
    for u in unavailable:
        ws.append([u["object"], u.get("link") or "", u["reason"]])
        r = ws.max_row
        link_cell(ws, r, 2, u.get("link"))
        ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True)
    style_sheet(ws, [40, 35, 70])

    out = path("competitors_vk_report.xlsx")
    wb.save(out)
    print(f"XLSX saved: {os.path.basename(out)} | sheets: {wb.sheetnames}")
    print("OUTPUTS DONE")
    return 0


if __name__ == "__main__":
    sys.exit(main())
